"""Data storage module for the Lyrics App."""
from pathlib import Path
import json
from fpdf import FPDF
from docx import Document
import os

# Define a base directory for storing data files
BASE_DIR = Path.home()  # This will use the user's home directory
APP_DIR = BASE_DIR / "RapWriter"  # Base directory for the app
LYRICS_DIR = APP_DIR / "lyrics"  # Subdirectory for lyrics
LOG_DIR = APP_DIR / "logs"  # Subdirectory for logs
SETTINGS_FILE = APP_DIR / "settings.json"

# Ensure the necessary directories exist
LYRICS_DIR.mkdir(parents=True, exist_ok=True)
LOG_DIR.mkdir(parents=True, exist_ok=True)


def save_lyrics(filename, lyrics):
    """Saves lyrics to a text file."""
    filepath = LYRICS_DIR / filename
    with filepath.open('w', encoding='utf-8') as file:
        file.write(lyrics)
    print(f"Lyrics saved to {filepath}")


def load_lyrics(filename):
    """Loads lyrics from a text file."""
    filepath = LYRICS_DIR / filename
    try:
        with filepath.open('r', encoding='utf-8') as file:
            lyrics = file.read()
        print(f"Lyrics loaded from {filepath}")
        return lyrics
    except FileNotFoundError:
        print(f"File {filepath} does not exist.")
        return ""


def save_settings(settings):
    """Saves application settings to a JSON file."""
    with SETTINGS_FILE.open('w', encoding='utf-8') as file:
        json.dump(settings, file, indent=4)
    print(f"Settings saved to {SETTINGS_FILE}")


def load_settings():
    """Loads application settings from a JSON file."""
    try:
        with SETTINGS_FILE.open('r', encoding='utf-8') as file:
            settings = json.load(file)
        print(f"Settings loaded from {SETTINGS_FILE}")
        return settings
    except FileNotFoundError:
        print(f"Settings file {SETTINGS_FILE} does not exist. "
              "Returning default settings.")
        return {}  # Return default settings if the file doesn't exist


def get_available_lyrics():
    """Returns a list of available lyrics files."""
    return [file.name for file in LYRICS_DIR.glob('*.txt')]


def export_to_pdf(filename, lyrics):
    """Exports lyrics to a PDF file."""
    filepath = LYRICS_DIR / filename.replace('.txt', '.pdf')
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for line in lyrics.split('\n'):
        pdf.cell(200, 10, txt=line, ln=True)

    # Create the directory if it doesn't exist
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    pdf.output(str(filepath))
    print(f"Lyrics exported to PDF at {filepath}")


def export_to_docx(filename, lyrics):
    """Exports lyrics to a DOCX file."""
    filepath = LYRICS_DIR / filename.replace('.txt', '.docx')
    doc = Document()
    doc.add_paragraph(lyrics)
    doc.save(filepath)
    print(f"Lyrics exported to DOCX at {filepath}")
